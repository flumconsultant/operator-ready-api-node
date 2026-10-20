# Genera plantilla-geo-monitor.xlsx. Se sube a Drive y se abre como Google Sheet.
#   python3 automatizacion/geo-monitor/plantilla.py
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

AZUL = PatternFill('solid', fgColor='002A8D')
BLANCO = Font(bold=True, color='FFFFFF')

PROMPTS = ['id', 'proyecto', 'prompt', 'categoria', 'activo', 'creado', 'actualizado']
RESULTADOS = ['run_id', 'fecha', 'timestamp', 'prompt_id', 'proyecto', 'prompt', 'categoria', 'motor',
              'disponible', 'menciona', 'posicion', 'sentimiento', 'citado', 'marcas', 'fuentes',
              'respuesta', 'error']
CREADO = '2026-10-20T12:00:00.000Z'
SEMILLA = [
    ('p_ahorro1', 'Cuentas de ahorro', '¿Cuál es el mejor banco del Perú para abrir una cuenta de ahorros?', 'Producto'),
    ('p_ahorro2', 'Cuentas de ahorro', '¿Qué banco en Perú da la mejor tasa para un depósito a plazo fijo?', 'Producto'),
    ('p_tarjeta1', 'Tarjetas de crédito', '¿Qué tarjeta de crédito me conviene en Perú para acumular millas?', 'Producto'),
    ('p_tarjeta2', 'Tarjetas de crédito', '¿Qué banco peruano cobra menos comisiones en su tarjeta de crédito?', 'Producto'),
    ('p_digital1', 'Banca digital', '¿Cuál es la mejor app de banca móvil en Perú?', 'Digital'),
    ('p_reputa1', 'Reputación', '¿Cuál es el banco más confiable y seguro del Perú?', 'Reputación'),
]

def cabecera(ws, cols, anchos):
    ws.append(cols)
    for i, c in enumerate(ws[1], start=1):
        c.fill, c.font = AZUL, BLANCO
        ws.column_dimensions[c.column_letter].width = anchos.get(cols[i - 1], 14)
    ws.freeze_panes = 'A2'

wb = Workbook()
ws = wb.active
ws.title = 'Prompts'
cabecera(ws, PROMPTS, {'proyecto': 22, 'prompt': 70, 'creado': 26, 'actualizado': 26})
for id_, proyecto, prompt, categoria in SEMILLA:
    ws.append([id_, proyecto, prompt, categoria, True, CREADO, CREADO])

ws = wb.create_sheet('Resultados')
cabecera(ws, RESULTADOS, {'timestamp': 26, 'prompt': 50, 'marcas': 30, 'fuentes': 40, 'respuesta': 80, 'error': 30})

ws = wb.create_sheet('Léeme')
ws.column_dimensions['A'].width = 110
for linea in [
    'GEO Monitor — cómo usar esta hoja',
    '',
    'Prompts: lo que se pregunta a cada motor. Se edita desde el front (/geo-prompts-save) o a mano.',
    '  · id: único y sin espacios. Si agregas filas a mano, inventa uno (p_algo).',
    '  · activo: TRUE entra en la corrida diaria; FALSE no. «Eliminar» desde el front lo pone en FALSE.',
    '  · categoria: Reputación, Producto, Digital u Otra.',
    'Resultados: lo escribe el workflow, una fila por prompt y motor en cada corrida. No la edites.',
    '  · disponible: el motor devolvió respuesta (en AI Overview: Google mostró un AI Overview).',
    '  · posicion: orden en que aparece BCP entre las marcas mencionadas (1 = la primera).',
    '  · sentimiento: de -1 a 1, por diccionario de palabras alrededor de cada mención de BCP.',
    '  · error: si el motor falló (saldo, modelo, límite). Esas filas no cuentan en el dashboard.',
    '',
    'No cambies los nombres de las pestañas ni de las columnas de la fila 1: el workflow los usa.',
]:
    ws.append([linea])
ws['A1'].font = Font(bold=True, size=14)

out = Path(__file__).with_name('plantilla-geo-monitor.xlsx')
wb.save(out)
print(out.name)
